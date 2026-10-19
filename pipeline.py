"""Slide deck -> narrated video.

Same recipe as the original Codex KVKK build:
  speaker notes -> ElevenLabs TTS (one clip per slide)
  Keynote (or LibreOffice) PDF export -> one still image per slide
  per slide: loudnorm -16 LUFS, 200 ms lead-in, >=600 ms tail, frame-exact still clip
  concat -> H.264/AAC MP4 with one chapter per slide, faststart
"""
import concurrent.futures
import hashlib
import json
import math
import os
import platform
import re
import shutil
import subprocess
import tempfile
import threading
import time
import wave
from pathlib import Path

import pymupdf
import imageio_ffmpeg
import requests
from pptx import Presentation

FF = imageio_ffmpeg.get_ffmpeg_exe()
API = 'https://api.elevenlabs.io'
FPS = 25
RATE = 48000
LEAD = 0.2           # silence before narration (seconds)
TAIL = 0.6           # minimum silence after narration (seconds)
SILENT_SLIDE = 3.0   # how long a slide without notes stays on screen
RESOLUTIONS = {'480p': (854, 480), '720p': (1280, 720), '1080p': (1920, 1080)}
DEFAULT_VOICE = '7VqWGAWwo2HMrylfKrcm'  # Fatih Yıldırım - Deep, Clear and Rich
DEFAULT_MODEL = 'eleven_v4'
FAKE_TTS = os.environ.get('SLIDE2VIDEO_FAKE_TTS') == '1'  # offline testing with macOS `say`
RENDERER = os.environ.get('SLIDE2VIDEO_RENDERER', '').strip().lower()  # 'keynote', 'libreoffice' or auto


class PipelineError(Exception):
    pass


def ffmpeg(args):
    r = subprocess.run([FF, '-hide_banner', '-loglevel', 'error', '-nostdin', '-y'] + args, capture_output=True)
    if r.returncode:
        raise PipelineError('ffmpeg hatası: ' + r.stderr.decode(errors='replace')[-800:])
    return r.stdout


# ---------------------------------------------------------------- slides

def read_deck(pptx_path):
    """Return every slide with its notes; hidden slides are flagged and later skipped."""
    deck = Presentation(str(pptx_path))
    slides = []
    for i, slide in enumerate(deck.slides, 1):
        notes = ''
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = ' '.join(slide.notes_slide.notes_text_frame.text.split())
        title = ''
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = ' '.join(slide.shapes.title.text_frame.text.split())
        if not title:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    title = ' '.join(shape.text_frame.text.split())
                    break
        slides.append({
            'slide': i,
            'title': title[:80] or f'Slayt {i}',
            'notes': notes,
            'chars': len(notes),
            'hidden': slide._element.get('show') == '0',
        })
    if not slides:
        raise PipelineError('Sunumda slayt yok.')
    return slides


EXPORT_SCRIPT = '''
on run argv
  set src to POSIX file (item 1 of argv)
  set dst to POSIX file (item 2 of argv)
  set docName to item 3 of argv
  with timeout of 900 seconds
    tell application id "com.apple.iWork.Keynote"
      launch
      -- Right after launch Keynote ignores scripting for a moment; wait until it answers.
      repeat 60 times
        try
          count of documents
          exit repeat
        on error
          delay 1
        end try
      end repeat
      open src
      -- open's return value is unreliable on a cold start, so find the document by its unique name.
      set docId to missing value
      repeat 120 times
        try
          set docId to id of (first document whose name starts with docName)
          exit repeat
        on error
          delay 1
        end try
      end repeat
      if docId is missing value then error "Keynote did not open the presentation."
      -- A freshly imported PPTX is sometimes not ready to export yet; retry by id.
      set lastErr to ""
      repeat with attempt from 1 to 6
        try
          if attempt < 6 then
            export (document id docId) to dst as PDF with properties {export style:IndividualSlides, skipped slides:true, all stages:false, PDF image quality:Best}
          else
            export (document id docId) to dst as PDF
          end if
          set lastErr to ""
          exit repeat
        on error e
          set lastErr to e
          delay 2
        end try
      end repeat
      try
        close (document id docId) saving no
      end try
      if lastErr is not "" then error lastErr
    end tell
  end timeout
end run
'''
_keynote = threading.Lock()  # one Keynote export at a time


def find_keynote():
    if platform.system() != 'Darwin':
        return None
    for app in (Path('/Applications/Keynote.app'), Path.home() / 'Applications/Keynote.app'):
        if app.exists():
            return app
    return None


def find_soffice():
    for name in ('soffice', 'libreoffice'):
        found = shutil.which(name)
        if found:
            return found
    candidates = [Path('/Applications/LibreOffice.app/Contents/MacOS/soffice'),
                  Path('/usr/bin/soffice'), Path('/usr/local/bin/soffice'), Path('/snap/bin/libreoffice')]
    candidates += sorted(Path('/opt').glob('libreoffice*/program/soffice'), reverse=True)
    for env in ('ProgramFiles', 'ProgramFiles(x86)', 'ProgramW6432'):
        if os.environ.get(env):
            candidates.append(Path(os.environ[env]) / 'LibreOffice/program/soffice.exe')
    for c in candidates:
        if c.exists():
            return str(c)
    return None


def renderer():
    """Keynote when available (truest to the original), otherwise LibreOffice."""
    if RENDERER not in ('', 'keynote', 'libreoffice'):
        raise PipelineError(f"SLIDE2VIDEO_RENDERER yalnızca 'keynote' veya 'libreoffice' olabilir, '{RENDERER}' verildi.")
    if RENDERER != 'libreoffice' and find_keynote():
        return 'keynote'
    if RENDERER == 'keynote':
        raise PipelineError('Keynote bulunamadı. Keynote yalnızca macOS\'ta çalışır; App Store\'dan ücretsiz yükleyebilirsiniz.')
    if find_soffice():
        return 'libreoffice'
    raise PipelineError('Slaytları dışa aktarmak için Keynote (yalnızca macOS) veya LibreOffice gerekli. '
                        'LibreOffice\'i https://www.libreoffice.org/download/ adresinden ücretsiz yükleyip tekrar deneyin.')


RENDERER_NAMES = {'keynote': 'Keynote', 'libreoffice': 'LibreOffice'}


def export_pdf(pptx_path, pdf_path):
    if pdf_path.exists():
        return
    if renderer() == 'libreoffice':
        export_pdf_libreoffice(pptx_path, pdf_path)
    else:
        export_pdf_keynote(pptx_path, pdf_path)


def export_pdf_libreoffice(pptx_path, pdf_path):
    """Headless LibreOffice; hidden slides are left out of the PDF (visible-only page count)."""
    soffice = find_soffice()
    with tempfile.TemporaryDirectory(dir=pdf_path.parent) as tmp:
        tmp = Path(tmp)
        # A private profile avoids clashing with an open LibreOffice window or a parallel export.
        profile = (tmp / 'profile').resolve().as_uri()
        try:
            r = subprocess.run([soffice, f'-env:UserInstallation={profile}', '--headless', '--norestore',
                                '--convert-to', 'pdf', '--outdir', str(tmp), str(pptx_path)],
                               capture_output=True, text=True, timeout=900)
        except subprocess.TimeoutExpired:
            raise PipelineError('LibreOffice dışa aktarımı 15 dakika içinde bitmedi.')
        out = tmp / (Path(pptx_path).stem + '.pdf')
        if r.returncode or not out.exists():
            raise PipelineError('LibreOffice dışa aktarımı başarısız oldu: ' + (r.stderr or r.stdout).strip()[-800:])
        shutil.move(str(out), str(pdf_path))


def export_pdf_keynote(pptx_path, pdf_path):
    """Keynote renders the deck with the real fonts/layout, like the 'Orijinal Slaytlar' build."""
    tmp = pdf_path.with_suffix('.part.pdf')
    # A uniquely named hard link lets the script find this exact document in Keynote.
    name = f'slide2video-{os.urandom(4).hex()}'
    link = pdf_path.parent / f'{name}.pptx'
    os.link(pptx_path, link)
    try:
        with _keynote:
            r = subprocess.run(['osascript', '-', str(link.resolve()), str(tmp.resolve()), name],
                               input=EXPORT_SCRIPT, text=True, capture_output=True, timeout=960)
    finally:
        link.unlink(missing_ok=True)
    if r.returncode or not tmp.exists():
        if '-1743' in r.stderr:
            raise PipelineError("Bu uygulamanın Keynote'u kontrol etme izni yok. Sistem Ayarları > Gizlilik ve "
                                "Güvenlik > Otomasyon bölümünden Terminal için Keynote iznini açıp tekrar deneyin.")
        raise PipelineError('Keynote dışa aktarımı başarısız oldu: ' + r.stderr.strip())
    tmp.replace(pdf_path)


def render_pages(pdf_path, out_dir, size):
    """Render each PDF page so it fits inside the target frame (sharp, no upscaling)."""
    out_dir.mkdir(exist_ok=True)
    w, h = size
    doc = pymupdf.open(str(pdf_path))
    paths = []
    for i, page in enumerate(doc, 1):
        p = out_dir / f'page-{i:03d}.png'
        if not p.exists():
            zoom = min(w / page.rect.width, h / page.rect.height)
            page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).save(str(p))
        paths.append(p)
    return paths


# ---------------------------------------------------------------- ElevenLabs

def _raise_for(r):
    try:
        detail = r.json().get('detail', r.text)
    except ValueError:
        detail = r.text
    if isinstance(detail, dict):
        detail = detail.get('message') or json.dumps(detail)
    if r.status_code == 401:
        raise PipelineError(f'ElevenLabs isteği reddetti (401), API anahtarını veya kredinizi kontrol edin: {detail}')
    raise PipelineError(f'ElevenLabs hatası {r.status_code}: {detail}')


def _post_audio(url, key, body):
    for attempt in range(5):
        r = requests.post(url, params={'output_format': 'mp3_44100_128'},
                          headers={'xi-api-key': key, 'Content-Type': 'application/json'},
                          json=body, timeout=300)
        if r.status_code == 200 and len(r.content) > 1000:
            return r.content
        if r.status_code in (429, 500, 502, 503, 504) and 'quota' not in r.text:
            time.sleep(2 ** attempt * 2)
            continue
        _raise_for(r)
    _raise_for(r)


def _dialogue_chunks(text, limit=1900):
    """Text-to-Dialogue accepts ~2000 chars per request; split on sentence ends."""
    chunks, cur = [], ''
    for sentence in text.replace('. ', '.\n').replace('? ', '?\n').replace('! ', '!\n').split('\n'):
        if cur and len(cur) + len(sentence) + 1 > limit:
            chunks.append(cur)
            cur = sentence
        else:
            cur = f'{cur} {sentence}'.strip()
    if cur:
        chunks.append(cur)
    return chunks


def synthesize(text, key, voice, model, language, cache_dir):
    """One narration clip, cached by (text, voice, model, language) so reruns never pay twice."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(json.dumps([text, voice, model, language, FAKE_TTS]).encode()).hexdigest()[:24]
    dest = cache_dir / f'{digest}.mp3'
    if dest.exists():
        return dest, True
    tmp = dest.with_suffix('.part')
    if FAKE_TTS:
        aiff = dest.with_suffix('.aiff')
        spoken = ' '.join(re.sub(r'\[[^\[\]]*\]', ' ', text).split())  # `say` would read tags aloud
        subprocess.run(['say', '-r', '190', '-o', str(aiff), spoken], check=True)
        ffmpeg(['-i', str(aiff), '-c:a', 'libmp3lame', '-b:a', '128k', '-f', 'mp3', str(tmp)])
        aiff.unlink()
    else:
        body = {'text': text, 'model_id': model}
        if language:
            body['language_code'] = language
        try:
            audio = _post_audio(f'{API}/v1/text-to-speech/{voice}', key, body)
        except PipelineError as e:
            # Some models (eleven_v4) may be served only through Text-to-Dialogue.
            if 'model' not in str(e).lower():
                raise
            parts = []
            for chunk in _dialogue_chunks(text):
                dbody = {'inputs': [{'text': chunk, 'voice_id': voice}], 'model_id': model}
                if language:
                    dbody['language_code'] = language
                parts.append(_post_audio(f'{API}/v1/text-to-dialogue', key, dbody))
            audio = b''.join(parts)  # MP3 frames concatenate cleanly
        tmp.write_bytes(audio)
    tmp.replace(dest)
    return dest, False


def list_voices(key):
    voices, token = [], None
    while True:
        params = {'page_size': 100}
        if token:
            params['next_page_token'] = token
        r = requests.get(f'{API}/v2/voices', headers={'xi-api-key': key}, params=params, timeout=30)
        if r.status_code != 200:
            _raise_for(r)
        data = r.json()
        voices += [{'voice_id': v['voice_id'], 'name': v.get('name') or v['voice_id'],
                    'category': v.get('category', '')} for v in data.get('voices', [])]
        token = data.get('next_page_token')
        if not data.get('has_more') or not token:
            return voices


def subscription(key):
    r = requests.get(f'{API}/v1/user/subscription', headers={'xi-api-key': key}, timeout=30)
    if r.status_code != 200:
        _raise_for(r)
    d = r.json()
    return {'used': d.get('character_count'), 'limit': d.get('character_limit'), 'tier': d.get('tier')}


# ---------------------------------------------------------------- video

def atempo_chain(speed):
    """atempo accepts 0.5-2.0 per stage; chain stages for anything outside that."""
    parts, s = [], speed
    while s > 2.0:
        parts.append('atempo=2.0')
        s /= 2.0
    while s < 0.5:
        parts.append('atempo=0.5')
        s /= 0.5
    parts.append(f'atempo={s:.6f}')
    return ','.join(parts)


def build_segment(n, image, audio, speed, size, seg_dir):
    """Still image + narration, both exactly `frames / FPS` seconds long."""
    w, h = size
    if audio is not None:
        filters = 'loudnorm=I=-16:TP=-1.5:LRA=11'
        if abs(speed - 1.0) > 1e-3:
            filters = atempo_chain(speed) + ',' + filters
        pcm = ffmpeg(['-i', str(audio), '-vn', '-af', filters, '-ac', '1', '-ar', str(RATE), '-f', 's16le', '-'])
        voice = len(pcm) / (RATE * 2)
        frames = math.ceil((voice + LEAD + TAIL) * FPS)
    else:
        pcm, voice = b'', 0.0
        frames = round(SILENT_SLIDE * FPS)
    seconds = frames / FPS
    lead = round(LEAD * RATE) if pcm else 0
    tail = round(seconds * RATE) - lead - len(pcm) // 2
    wav = seg_dir / f'seg-{n:03d}.wav'
    with wave.open(str(wav), 'wb') as f:
        f.setnchannels(1)
        f.setsampwidth(2)
        f.setframerate(RATE)
        f.writeframes(b'\0' * (lead * 2) + pcm + b'\0' * (tail * 2))
    video = seg_dir / f'seg-{n:03d}.mp4'
    vf = (f'scale={w}:{h}:force_original_aspect_ratio=decrease:flags=lanczos,'
          f'pad={w}:{h}:(ow-iw)/2:(oh-ih)/2:color=white,setsar=1')
    ffmpeg(['-loop', '1', '-framerate', str(FPS), '-i', str(image), '-vf', vf, '-frames:v', str(frames),
            '-an', '-c:v', 'libx264', '-preset', 'fast', '-tune', 'stillimage', '-crf', '18',
            '-pix_fmt', 'yuv420p', '-r', str(FPS), '-video_track_timescale', '12800', '-threads', '2', str(video)])
    return {'slide': n, 'seconds': seconds, 'voice_seconds': round(voice, 2), 'video': video, 'audio': wav}


def concat(segments, titles, out_path, work, title):
    meta = [';FFMETADATA1', f'title={title}']
    cursor = 0.0
    for seg in segments:
        end = cursor + seg['seconds']
        meta += ['[CHAPTER]', 'TIMEBASE=1/1000', f'START={round(cursor * 1000)}', f'END={round(end * 1000)}',
                 'title=' + titles[seg['slide']].replace('=', '-').replace(';', ',').replace('#', '')]
        cursor = end
    (work / 'chapters.txt').write_text('\n'.join(meta) + '\n')
    (work / 'video-list.txt').write_text(''.join(f"file '{s['video']}'\n" for s in segments))
    (work / 'audio-list.txt').write_text(''.join(f"file '{s['audio']}'\n" for s in segments))
    ffmpeg(['-f', 'concat', '-safe', '0', '-i', str(work / 'video-list.txt'),
            '-f', 'concat', '-safe', '0', '-i', str(work / 'audio-list.txt'),
            '-i', str(work / 'chapters.txt'), '-map', '0:v:0', '-map', '1:a:0',
            '-map_metadata', '2', '-map_chapters', '2', '-c:v', 'copy', '-c:a', 'aac', '-b:a', '128k',
            '-movflags', '+faststart', str(out_path)])
    return cursor


# ---------------------------------------------------------------- whole job

def tag_summary(info):
    msg = ('Etiketleri yapay zekâ yerleştirdi (' + info['method'] + ')') if info['method'] != 'kural' \
        else 'Etiketler kural tabanlı eklendi'
    def nums(ns, many, one):
        return f'{", ".join(map(str, ns))} numaralı {many if len(ns) > 1 else one}'
    if info['fallback']:
        msg += f'; {nums(info["fallback"], "slaytlarda", "slaytta")} model sonucu kullanılamadığı için kurallar uygulandı'
    if info['kept']:
        msg += f'; {nums(info["kept"], "slaytlardaki", "slayttaki")} mevcut etiketlere dokunulmadı'
    return msg + '.'


def run(job_dir, pptx_path, opts, cache_dir, report):
    """opts: api_key, voice_id, model, language, speed, resolution, narrations / auto_tags + llm_key.
    report(stage, progress, message)."""
    size = RESOLUTIONS[opts['resolution']]
    speed = float(opts['speed'])
    free = shutil.disk_usage(job_dir).free
    if free < 600 * 1024 * 1024:
        raise PipelineError(f'Diskte yalnızca {free // 2**20} MB boş alan var; en az 600 MB boşaltıp tekrar deneyin.')

    slides = [s for s in read_deck(pptx_path) if not s['hidden']]
    all_count = len(read_deck(pptx_path))

    tool = RENDERER_NAMES[renderer()]
    report(f'Slaytlar {tool} ile dışa aktarılıyor', 0.02, f'{len(slides)} görünür slayt')
    pdf = job_dir / 'slides.pdf'
    export_pdf(pptx_path, pdf)
    pages = render_pages(pdf, job_dir / f'pages-{opts["resolution"]}', size)
    if len(pages) == all_count:
        images = {s['slide']: pages[s['slide'] - 1] for s in slides}
    elif len(pages) == len(slides):
        images = {s['slide']: p for s, p in zip(slides, pages)}
    else:
        raise PipelineError(f'{tool} {all_count} slayt için {len(pages)} sayfa üretti '
                            f'({len(slides)} görünür slayt bekleniyordu).')

    narrations = {int(k): v for k, v in (opts.get('narrations') or {}).items()}
    if opts.get('auto_tags') and not narrations:
        import tagging
        report('İfade etiketleri ekleniyor', 0.05, 'Notlara etiket ekleniyor')
        narrations, info = tagging.tag_notes(slides, opts.get('llm_key', ''))
        report('İfade etiketleri ekleniyor', 0.07, tag_summary(info))
    for s in slides:
        if s['notes'] and (narrations.get(s['slide']) or '').strip():
            s['notes'] = ' '.join(narrations[s['slide']].split())
            s['chars'] = len(s['notes'])

    voiced = [s for s in slides if s['notes']]
    audio, done = {}, 0
    report('Seslendirme oluşturuluyor', 0.08, f'Notlu {len(voiced)} slayt, {sum(s["chars"] for s in voiced)} karakter')

    def tts(s):
        return s['slide'], synthesize(s['notes'], opts['api_key'], opts['voice_id'], opts['model'],
                                      opts['language'], cache_dir)
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as ex:
        futures = [ex.submit(tts, s) for s in voiced]
        try:
            for fut in concurrent.futures.as_completed(futures):
                n, (path, cached) = fut.result()
                audio[n] = path
                done += 1
                report('Seslendirme oluşturuluyor', 0.08 + 0.52 * done / max(len(voiced), 1),
                       f'Slayt {n}: {"kayıtlı ses tekrar kullanıldı" if cached else "seslendirildi"}')
        except Exception:
            for f in futures:
                f.cancel()
            raise

    seg_dir = job_dir / 'segments'
    shutil.rmtree(seg_dir, ignore_errors=True)
    seg_dir.mkdir()
    segments, done = {}, 0
    report('Video oluşturuluyor', 0.62, 'Slayt klipleri kodlanıyor')
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as ex:
        futures = [ex.submit(build_segment, s['slide'], images[s['slide']], audio.get(s['slide']), speed, size, seg_dir)
                   for s in slides]
        for fut in concurrent.futures.as_completed(futures):
            seg = fut.result()
            segments[seg['slide']] = seg
            done += 1
            report('Video oluşturuluyor', 0.62 + 0.3 * done / len(slides), f'Slayt {seg["slide"]} kodlandı ({seg["seconds"]:.1f} sn)')

    ordered = [segments[s['slide']] for s in slides]
    titles = {s['slide']: f'{s["slide"]:02d} · {s["title"]}' for s in slides}
    name = Path(opts.get('original_name') or 'slides').stem
    tag = f'{opts["resolution"]}' + ('' if abs(speed - 1) < 1e-3 else f'_{speed:g}x')
    out = job_dir / f'{name}_{tag}.mp4'
    report('Son video birleştiriliyor', 0.94, 'Klipler birleştiriliyor, bölümler ekleniyor')
    total = concat(ordered, titles, out, job_dir, name)
    shutil.rmtree(seg_dir, ignore_errors=True)  # disk is tight; keep only the result
    ffmpeg(['-xerror', '-i', str(out), '-f', 'null', '-'])  # full decode check
    return {'file': out.name, 'seconds': round(total, 2), 'bytes': out.stat().st_size, 'slides': len(slides)}
