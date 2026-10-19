"""Smoke-test a built package: start it, wait for the server, upload a small deck, stop it.

    python packaging/smoke.py <path to SlayttanVideoya.exe or the .app's executable>
"""
import json
import os
import subprocess
import sys
import tempfile
import time
import urllib.request
import uuid
from pathlib import Path

from pptx import Presentation

URL = 'http://127.0.0.1:8765'


def main():
    exe = sys.argv[1]
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        deck = tmp / 'deneme.pptx'
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = 'Deneme'
        s.notes_slide.notes_text_frame.text = 'Bu bir deneme notudur.'
        p.save(deck)

        proc = subprocess.Popen([exe], env={**os.environ, 'SLIDE2VIDEO_DATA': str(tmp / 'data')},
                                stdin=subprocess.DEVNULL)
        try:
            for _ in range(60):
                try:
                    cfg = json.load(urllib.request.urlopen(URL + '/api/config', timeout=2))
                    break
                except OSError:
                    if proc.poll() is not None:
                        sys.exit(f'Uygulama kapandı (çıkış kodu {proc.returncode}).')
                    time.sleep(1)
            else:
                sys.exit('Sunucu 60 saniyede açılmadı.')
            print('config:', cfg)
            assert urllib.request.urlopen(URL + '/').status == 200

            boundary = uuid.uuid4().hex
            body = (f'--{boundary}\r\nContent-Disposition: form-data; name="deck"; filename="deneme.pptx"\r\n'
                    'Content-Type: application/octet-stream\r\n\r\n').encode() + deck.read_bytes() \
                + f'\r\n--{boundary}--\r\n'.encode()
            req = urllib.request.Request(URL + '/api/upload', data=body,
                                         headers={'Content-Type': f'multipart/form-data; boundary={boundary}'})
            job = json.load(urllib.request.urlopen(req))
            assert job['slides'][0]['notes'] == 'Bu bir deneme notudur.', job
            assert (tmp / 'data/jobs' / job['id'] / 'deck.pptx').exists(), 'data klasörü kullanılmadı'
            print('upload ok:', job['id'])
        finally:
            proc.kill()
            proc.wait()


if __name__ == '__main__':
    main()
